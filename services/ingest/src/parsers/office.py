#!/usr/bin/env python3
"""办公扩展解析器（08 §3.3 的 Word / Excel / PPT / PDF 文本层）。

由 ``services/ingest`` 的 TypeScript 包装在办公扩展的解释器里跑。
**不出网**：脚本里没有网络客户端；强制点仍在宿主的受限子进程。

stdout 不写正文。结果写到 ``--result`` 指定的 JSON 文件，形状与 TS 侧 ``ParseResult`` 对齐。
报错只给机器码，不回显文件内容（Q14）。
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable

PARSER_VERSION = "1"
MARKDOWN_ROW_LIMIT_DEFAULT = 200

EXIT_OK = 0
EXIT_UNPARSED = 2
EXIT_RUNTIME_MISSING = 3
EXIT_DAMAGED = 4


def fail(code: int, reason: str) -> None:
    sys.stderr.write(json.dumps({"ok": False, "code": reason}, ensure_ascii=True) + "\n")
    raise SystemExit(code)


def require(name: str):
    try:
        return __import__(name)
    except ImportError:
        fail(EXIT_RUNTIME_MISSING, "RUNTIME_MISSING")


def write_result(path: Path, markdown: str, meta: dict[str, Any], assets: list[str]) -> None:
    if markdown.strip() == "" and not assets:
        fail(EXIT_UNPARSED, "EMPTY")
    path.write_text(
        json.dumps(
            {"markdown": markdown, "meta": meta, "assets": assets},
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


def meta_of(
    parser: str,
    markdown: str,
    *,
    tables: int,
    pages: int | None = None,
    confidence: float = 1.0,
    partial: bool = False,
    note: str | None = None,
) -> dict[str, Any]:
    out: dict[str, Any] = {
        "parser": parser,
        "parserVersion": PARSER_VERSION,
        "chars": len(markdown),
        "tables": tables,
        "confidence": confidence,
    }
    if pages is not None:
        out["pages"] = pages
    if partial:
        out["partial"] = True
    if note:
        out["note"] = note
    return out


def escape_cell(value: str) -> str:
    return value.replace("\n", " ").replace("|", "\\|").strip()


def table_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    body = normalized[1:] if len(normalized) > 1 else []
    lines = [
        "| " + " | ".join(escape_cell(c) for c in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
        *[
            "| " + " | ".join(escape_cell(c) for c in row) + " |"
            for row in body
        ],
    ]
    return "\n".join(lines)


def save_blob(assets_dir: Path, stem: str, blob: bytes, content_type: str | None) -> str:
    assets_dir.mkdir(parents=True, exist_ok=True)
    ext = "bin"
    if content_type:
        subtype = content_type.split("/")[-1].lower()
        ext = {"jpeg": "jpg", "svg+xml": "svg"}.get(subtype, subtype)
    elif blob[:8] == b"\x89PNG\r\n\x1a\n":
        ext = "png"
    elif blob[:2] == b"\xff\xd8":
        ext = "jpg"
    name = f"{stem}.{ext}"
    (assets_dir / name).write_bytes(blob)
    return f"assets/{name}"


def heading_level(paragraph) -> int | None:
    style = paragraph.style
    name = style.name if style is not None else ""
    name = name or ""
    for prefix in ("Heading ", "标题 "):
        if name.startswith(prefix):
            try:
                return max(1, min(6, int(name[len(prefix) :].strip())))
            except ValueError:
                break
    try:
        pPr = paragraph._p.pPr
        if pPr is not None and pPr.outlineLvl is not None:
            val = pPr.outlineLvl.val
            if val is not None:
                return max(1, min(6, int(val) + 1))
    except Exception:
        pass
    return None


def is_list_item(paragraph) -> bool:
    name = ""
    if paragraph.style is not None:
        name = paragraph.style.name or ""
    if "List" in name or "Bullet" in name or name.startswith("列表"):
        return True
    try:
        pPr = paragraph._p.pPr
        return pPr is not None and pPr.numPr is not None
    except Exception:
        return False


def iter_docx_blocks(document) -> Iterable[Any]:
    from docx.document import Document as DocumentClass
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body if isinstance(document, DocumentClass) else document._tc
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def parse_docx(path: Path, assets_dir: Path) -> tuple[str, dict[str, Any], list[str]]:
    Document = require("docx").Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(str(path))
    except Exception:
        fail(EXIT_DAMAGED, "DAMAGED")

    lines: list[str] = []
    tables = 0
    assets: list[str] = []
    image_index = 0

    def paragraph_md(paragraph) -> str | None:
        text = paragraph.text.strip()
        if text == "":
            return None
        level = heading_level(paragraph)
        if level is not None:
            return f"{'#' * level} {text}"
        if is_list_item(paragraph):
            return f"- {text}"
        return text

    for block in iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            rendered = paragraph_md(block)
            if rendered:
                lines.append(rendered)
            continue
        if not isinstance(block, Table):
            continue
        rows = [[cell.text for cell in row.cells] for row in block.rows]
        if len(rows) == 1 and len(rows[0]) == 1:
            title = rows[0][0].strip()
            if title:
                first, *rest = title.split("\n")
                lines.append(f"## {first.strip()}")
                lines.extend(part.strip() for part in rest if part.strip())
            continue
        tables += 1
        rendered = table_markdown(rows)
        if rendered:
            lines.append("")
            lines.append(rendered)
            lines.append("")

    try:
        for rel in document.part.rels.values():
            if "image" not in getattr(rel, "reltype", ""):
                continue
            image_index += 1
            part = rel.target_part
            assets.append(
                save_blob(
                    assets_dir,
                    f"img-{image_index}",
                    part.blob,
                    getattr(part, "content_type", None),
                )
            )
    except Exception:
        pass

    markdown = "\n".join(lines).strip() + ("\n" if lines else "")
    if assets:
        markdown += "\n" + "\n".join(f"![嵌入图]({item})" for item in assets) + "\n"
    return markdown, meta_of("office-docx", markdown, tables=tables), assets


def parse_xlsx(path: Path, assets_dir: Path, row_limit: int) -> tuple[str, dict[str, Any], list[str]]:
    openpyxl = require("openpyxl")
    try:
        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        formulas = openpyxl.load_workbook(str(path), data_only=False, read_only=True)
    except Exception:
        fail(EXIT_DAMAGED, "DAMAGED")

    assets_dir.mkdir(parents=True, exist_ok=True)
    lines: list[str] = []
    assets: list[str] = []
    tables = 0
    formula_sheets = {sheet.title: sheet for sheet in formulas}

    def cell_text(value: Any) -> str:
        if value is None:
            return ""
        return str(value)

    for sheet in workbook:
        tables += 1
        rows: list[list[str]] = []
        formula_rows: list[list[Any]] = []
        formula_sheet = formula_sheets.get(sheet.title)
        if formula_sheet is not None:
            formula_rows = [list(row) for row in formula_sheet.iter_rows(values_only=True)]
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            values: list[str] = []
            for column, cell in enumerate(row):
                if cell is not None:
                    values.append(cell_text(cell))
                    continue
                fallback = ""
                if index < len(formula_rows) and column < len(formula_rows[index]):
                    fallback = cell_text(formula_rows[index][column])
                values.append(fallback)
            if all(item == "" for item in values) and index > 0:
                continue
            rows.append(values)
        csv_name = re.sub(r"[^\w\u4e00-\u9fff-]+", "-", sheet.title).strip("-") or f"sheet-{tables}"
        csv_rel = f"assets/{csv_name}.csv"
        with (assets_dir / f"{csv_name}.csv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerows(rows)
        assets.append(csv_rel)
        shown = rows[: row_limit + 1] if rows else rows
        lines.append(f"## {sheet.title}")
        lines.append("")
        rendered = table_markdown(shown)
        if rendered:
            lines.append(rendered)
        if len(rows) > row_limit + 1:
            lines.append("")
            lines.append(
                f"> 只显示前 {row_limit} 行，共 {max(0, len(rows) - 1)} 行。"
                f"完整数据在 {csv_rel}。"
            )
        lines.append("")

    workbook.close()
    formulas.close()
    markdown = "\n".join(lines).strip() + "\n"
    note = "每个工作表另存了完整 csv，Markdown 里只放预览。"
    return markdown, meta_of("office-xlsx", markdown, tables=tables, note=note), assets


def parse_pptx(path: Path, assets_dir: Path) -> tuple[str, dict[str, Any], list[str]]:
    Presentation = require("pptx").Presentation
    try:
        deck = Presentation(str(path))
    except Exception:
        fail(EXIT_DAMAGED, "DAMAGED")

    lines: list[str] = []
    assets: list[str] = []
    image_index = 0
    for index, slide in enumerate(deck.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                )
                if text:
                    chunks.append(text)
            image = getattr(shape, "image", None)
            if image is not None:
                image_index += 1
                rel = save_blob(
                    assets_dir,
                    f"slide-{index}-img-{image_index}",
                    image.blob,
                    image.content_type,
                )
                assets.append(rel)
                chunks.append(f"![嵌入图]({rel})")
        heading = chunks[0].split("\n", 1)[0] if chunks else f"第 {index} 页"
        body: list[str] = []
        if chunks:
            body.extend(chunks[0].split("\n")[1:])
            body.extend(chunks[1:])
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        lines.append(f"## {heading}")
        lines.extend(part for part in body if part)
        if notes:
            lines.append(f"> 备注：{notes}")
        lines.append("")

    markdown = "\n".join(lines).strip() + "\n"
    note = "幻灯片未渲染成图片，只提取了文字、备注和嵌入图。"
    return (
        markdown,
        meta_of("office-pptx", markdown, tables=0, pages=len(deck.slides), note=note),
        assets,
    )


def parse_pdf(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    pdfplumber = require("pdfplumber")
    try:
        pdf = pdfplumber.open(str(path))
    except Exception as err:
        message = str(err).lower()
        if "password" in message or "encrypt" in message:
            fail(EXIT_DAMAGED, "ENCRYPTED")
        fail(EXIT_DAMAGED, "DAMAGED")

    lines: list[str] = []
    tables = 0
    with pdf:
        pages = len(pdf.pages)
        for index, page in enumerate(pdf.pages, start=1):
            lines.append(f"## 第 {index} 页")
            text = (page.extract_text() or "").strip()
            if text:
                lines.append(text)
            extracted = page.extract_tables() or []
            for table in extracted:
                rows = [["" if cell is None else str(cell) for cell in row] for row in table]
                if not rows:
                    continue
                tables += 1
                lines.append("")
                lines.append(table_markdown(rows))
                lines.append("")
            lines.append("")

    markdown = "\n".join(lines).strip() + "\n"
    if markdown.strip() == "":
        fail(EXIT_UNPARSED, "EMPTY")
    return markdown, meta_of("office-pdf", markdown, tables=tables, pages=pages), []


def parse_rtf(path: Path) -> tuple[str, dict[str, Any], list[str]]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin1")
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = re.sub(r"[{}]", " ", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    markdown = text + ("\n" if text else "")
    return markdown, meta_of("office-rtf", markdown, tables=0), []


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--kind", required=True)
    parser.add_argument("--input", required=True)
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--result", required=True)
    parser.add_argument("--row-limit", type=int, default=MARKDOWN_ROW_LIMIT_DEFAULT)
    args = parser.parse_args(argv)

    source = Path(args.input)
    out_dir = Path(args.out_dir)
    assets_dir = out_dir / "assets"
    if not source.is_file():
        fail(EXIT_DAMAGED, "MISSING_INPUT")

    head = source.read_bytes()[:8]
    kind = args.kind
    if kind == "docx" and head.startswith(b"\xd0\xcf\x11\xe0"):
        fail(EXIT_UNPARSED, "LEGACY_DOC")

    if kind == "docx":
        markdown, meta, assets = parse_docx(source, assets_dir)
    elif kind == "xlsx":
        markdown, meta, assets = parse_xlsx(source, assets_dir, args.row_limit)
    elif kind == "pptx":
        markdown, meta, assets = parse_pptx(source, assets_dir)
    elif kind == "pdf":
        markdown, meta, assets = parse_pdf(source)
    elif kind == "rtf":
        markdown, meta, assets = parse_rtf(source)
    else:
        fail(EXIT_UNPARSED, "UNSUPPORTED")

    write_result(Path(args.result), markdown, meta, assets)
    raise SystemExit(EXIT_OK)


if __name__ == "__main__":
    main()
