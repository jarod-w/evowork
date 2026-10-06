"""Bounded body extraction, separate from task summaries/Office previews. No asset export."""
import argparse
import json
import mimetypes
import resource
import zipfile
from pathlib import Path

LIMIT = 2_000_000
MAX_PAGES = 100
# No dependency may load host MIME configuration outside the approved parser scope.
mimetypes.knownfiles = []
mimetypes.init(files=[])


def extract(path):
    blocks, used, partial, pending = [], 0, False, 0

    def add(text, location, page=None, source="text"):
        nonlocal used, partial
        if not text.strip():
            return
        remaining = LIMIT - used
        if len(text) > remaining:
            partial = True
        text = text[:remaining]
        if text:
            block = {"text": text, "location": location, "source": source}
            if page is not None:
                block["page"] = page
            blocks.append(block)
            used += len(text)

    ext = path.suffix.lower()
    if ext in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(x.file_size for x in entries) > 512 * 1024 * 1024:
                raise ValueError("ARCHIVE_LIMIT")
            if any(x.file_size > 64 * 1024 * 1024 for x in entries):
                raise ValueError("ARCHIVE_LIMIT")
    if ext == ".pdf":
        import pypdfium2 as pdfium
        document = pdfium.PdfDocument(path)
        if pdfium.raw.FPDF_GetSecurityHandlerRevision(document) != -1:
            raise ValueError("PDF_ENCRYPTED")
        total = len(document)
        partial = total > MAX_PAGES
        for index in range(min(total, MAX_PAGES)):
            page = document[index]
            textpage = page.get_textpage()
            text = textpage.get_text_range()
            objects = list(page.get_objects())
            images = [obj for obj in objects if obj.type == pdfium.raw.FPDF_PAGEOBJ_IMAGE]
            # Mirrors the conservative OCR classifier: a small header does not prove a native page.
            width, height = page.get_size()
            area = max(1, width * height)
            coverage = 0
            for image in images:
                left, bottom, right, top = image.get_bounds()
                coverage += max(0, right-left) * max(0, top-bottom) / area
            suspect = sum(c == '\ufffd' or (ord(c) < 32 and c not in '\n\r\t') for c in text) > max(1,len(text)) * .1
            candidate = bool(objects) and (suspect or len(''.join(text.split())) < 32 or coverage >= .6)
            if candidate:
                pending += 1
                partial = True
            add(text, f"第 {index+1} 页", index+1, "textLayer")
            textpage.close()
            page.close()
            if used >= LIMIT:
                partial = index+1 < total or partial
                break
        document.close()
    elif ext == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        cells = 0
        for sheet in book:
            group, first = [], 1
            for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                cells += len(row)
                if cells > 2_000_000:
                    partial = True
                    break
                group.append(' | '.join(str(cell) if cell is not None else '' for cell in row))
                if index % 20 == 0:
                    add('\n'.join(group), f"工作表 {sheet.title} · 行 {first}–{index}")
                    group, first = [], index+1
                if used >= LIMIT:
                    partial = True
                    break
            if group:
                add('\n'.join(group), f"工作表 {sheet.title} · 行 {first}–{index}")
            if partial:
                break
        book.close()
    elif ext == ".docx":
        from docx import Document
        doc = Document(path)
        for index, paragraph in enumerate(doc.paragraphs, 1):
            add(paragraph.text, f"段落 {index}")
            if used >= LIMIT:
                partial = True
                break
        for index, table in enumerate(doc.tables, 1):
            add('\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows), f"表格 {index}")
            if used >= LIMIT:
                partial = True
                break
    elif ext == ".pptx":
        from pptx import Presentation
        deck = Presentation(path)
        for index, slide in enumerate(deck.slides, 1):
            text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text.append(shape.text)
                if shape.has_table:
                    text.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
            if slide.has_notes_slide:
                text.append(slide.notes_slide.notes_text_frame.text)
            add('\n'.join(text), f"幻灯片 {index}", index)
            if used >= LIMIT:
                partial = True
                break
    else:
        raise ValueError("UNSUPPORTED")
    return {"blocks": blocks, "partial": partial, "ocrCandidates": pending}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True)
    parser.add_argument('--result', required=True)
    args = parser.parse_args()
    resource.setrlimit(resource.RLIMIT_CPU, (30, 30))
    resource.setrlimit(resource.RLIMIT_FSIZE, (32 * 1024 * 1024, 32 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
    try:
        result = extract(Path(args.input))
        Path(args.result).write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    except Exception:
        raise SystemExit(2)


if __name__ == '__main__':
    main()
