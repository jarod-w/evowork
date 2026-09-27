"""
多附件 E2E 的输入：一份带两张图的 docx、一份带两张图的 pptx、一份 xlsx、一份 csv、一张 png。

**每个文件都藏一个只有它才有的事实**，E2E 让模型逐条答出来：
  · docx 正文：上半年累计 915 万元
  · pptx 饼图：企业客户 45% —— **这个数只在图的像素里**，幻灯片文字里没有。
    答得出它，才证明文档里嵌的图真的抽出来、真的作为 localImage 进了模型
  · xlsx：Q3 预算三项合计 80 万元（只在工作表里）
  · csv：两个渠道（星河科技 · 云帆数据）
  · png：图例里的两条线「销售额」「成本」

用法：<办公扩展的 python> make-office-fixtures.py <输出目录>
"""
import sys, io
from pathlib import Path
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PI, Pt
from openpyxl import Workbook

out = Path(sys.argv[1]); out.mkdir(exist_ok=True)
fonts = list((Path.home()/".evowork/runtime/office/fonts").glob("*"))
for f in fonts: font_manager.fontManager.addfont(str(f))
if fonts: plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(fonts[0])).get_name()

months = ["1月","2月","3月","4月","5月","6月"]
sales = [120, 135, 150, 142, 170, 198]
cost  = [80, 88, 95, 99, 104, 118]

def chart(kind, name):
    fig, ax = plt.subplots(figsize=(6,3.5), dpi=100)
    if kind == "bar":
        ax.bar(months, sales, color="#2f6fed"); ax.set_title("华东区月度销售额（万元）")
    elif kind == "line":
        ax.plot(months, sales, marker="o", label="销售额"); ax.plot(months, cost, marker="s", label="成本"); ax.legend(); ax.set_title("销售额与成本走势")
    else:
        ax.pie([45,30,25], labels=["企业客户","中小客户","个人"], autopct="%d%%"); ax.set_title("客户结构")
    p = out/f"{name}.png"; fig.tight_layout(); fig.savefig(p); plt.close(fig); return p

bar, line, pie = chart("bar","bar"), chart("line","line"), chart("pie","pie")

d = Document()
d.add_heading("2026 上半年华东区销售报告", 0)
d.add_paragraph("上半年华东区销售额累计 915 万元，同比增长 18%。6 月单月达到 198 万元，创年内新高。")
d.add_heading("一、月度销售", 1)
d.add_picture(str(bar), width=Inches(5.5))
d.add_paragraph("图 1 显示销售额在 4 月小幅回落后连续两个月增长。")
d.add_heading("二、成本", 1)
d.add_picture(str(line), width=Inches(5.5))
t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
for c, h in zip(t.rows[0].cells, ["月份","销售额","成本"]): c.text = h
for m, s, c in zip(months, sales, cost):
    r = t.add_row().cells; r[0].text, r[1].text, r[2].text = m, str(s), str(c)
d.add_paragraph("结论：毛利率从 33% 提升到 40%，主要来自企业客户的放量。")
d.save(out/"上半年销售报告.docx")

pr = Presentation()
s = pr.slides.add_slide(pr.slide_layouts[0]); s.shapes.title.text = "Q3 市场计划"; s.placeholders[1].text = "市场部 · 2026-07"
s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = "客户结构"
s.shapes.add_picture(str(pie), PI(1.5), PI(1.6), width=PI(7))
s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = "Q3 目标"
tb = s.shapes.add_textbox(PI(1), PI(1.8), PI(8), PI(3)).text_frame
tb.text = "销售额目标：Q3 达到 650 万元"
for line_ in ["企业客户占比提升到 50%", "新增 3 个渠道伙伴", "市场预算明细见预算表"]:
    p = tb.add_paragraph(); p.text = line_; p.font.size = Pt(20)
s.shapes.add_picture(str(bar), PI(5.5), PI(4.2), width=PI(4))
pr.save(out/"Q3市场计划.pptx")

wb = Workbook(); ws = wb.active; ws.title = "预算"
ws.append(["项目","Q3 预算（万元）"])
for row in [["线上广告",35],["线下活动",25],["渠道激励",20]]: ws.append(row)
wb.save(out/"市场预算.xlsx")

(out/"渠道名单.csv").write_text("渠道,城市,签约月份\n星河科技,杭州,7月\n云帆数据,苏州,8月\n", encoding="utf-8")
import shutil; shutil.copy(line, out/"走势图.png")
for p in (bar, line, pie): p.unlink()
print("\n".join(sorted(x.name for x in out.iterdir())))
