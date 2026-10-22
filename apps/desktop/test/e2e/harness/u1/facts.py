"""读一份产物，列出**机器能确定的事实**，给 U1 人工评分做参考（08 §5.4）。

只列事实、不打分：「文字不溢出、中文不乱码」要看渲染结果，这里判不了；
这里给的是评分人一眼核对不了、却影响判断的东西 —— 页数、表格与目录在不在、
公式有几条、有没有写死的数、用了哪些字体（东亚字体缺失是乱码的头号原因）。

用法：python3 facts.py <产物路径>   → 一行 JSON
"""

import json
import re
import sys
from pathlib import Path

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def pptx_facts(path: Path) -> dict:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(path))
    fonts, east_asian = set(), set()
    with_text, charts, pictures, tables = 0, 0, 0, 0
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for slide in prs.slides:
        # 「有文字的页」而不是「有标题占位符的页」：随包技能用普通文本框排版，按占位符数会一律是 0
        if any(shape.has_text_frame and shape.text_frame.text.strip() for shape in slide.shapes):
            with_text += 1
        for shape in slide.shapes:
            charts += int(getattr(shape, "has_chart", False))
            tables += int(getattr(shape, "has_table", False))
            pictures += int(shape.shape_type == 13)
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                # 字体可能写在段落的默认属性上（随包技能就是这么写的），文字片段上没有
                if paragraph.font.name:
                    fonts.add(paragraph.font.name)
                for node in paragraph._p.iter(f"{ns}ea"):
                    if node.get("typeface"):
                        east_asian.add(node.get("typeface"))
                for run in paragraph.runs:
                    if run.font.name:
                        fonts.add(run.font.name)
    return {
        "slides": len(prs.slides),
        "slidesWithText": with_text,
        "charts": charts,
        "tables": tables,
        "pictures": pictures,
        "fonts": sorted(fonts),
        "eastAsianFonts": sorted(east_asian),
        "slideSize": [round(Emu(prs.slide_width).inches, 2), round(Emu(prs.slide_height).inches, 2)],
    }


def heading_level(paragraph) -> int | None:
    """标题级别：样式名（Heading N / 标题 N）或大纲级别（随包技能用的是「正文 + 大纲级别」）。"""
    name = paragraph.style.name if paragraph.style is not None else ""
    match = re.match(r"(Heading|标题)\s*(\d)", name)
    if match:
        return int(match.group(2))
    for element in (paragraph._p, paragraph.style.element if paragraph.style is not None else None):
        node = element.find(f".//{W_NS}outlineLvl") if element is not None else None
        if node is not None and node.get(f"{W_NS}val", "").isdigit() and int(node.get(f"{W_NS}val")) < 9:
            return int(node.get(f"{W_NS}val")) + 1
    return None


def docx_facts(path: Path) -> dict:
    from docx import Document

    doc = Document(str(path))
    levels: dict[str, int] = {}
    for paragraph in doc.paragraphs:
        level = heading_level(paragraph)
        if level is not None:
            levels[f"H{level}"] = levels.get(f"H{level}", 0) + 1
    body = doc.element.body
    instr = " ".join(node.text or "" for node in body.iter(f"{W_NS}instrText"))
    has_sdt_toc = any("Table of Contents" in (n.get(f"{W_NS}val") or "") for n in body.iter(f"{W_NS}docPartGallery"))
    section = doc.sections[0]
    fonts = set()
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name:
                fonts.add(run.font.name)
            rfonts = run._r.find(f".//{W_NS}rFonts")
            if rfonts is not None and rfonts.get(f"{W_NS}eastAsia"):
                fonts.add(rfonts.get(f"{W_NS}eastAsia"))
    return {
        "headings": levels,
        "tables": len(doc.tables),
        "toc": "TOC" in instr or has_sdt_toc,
        "paragraphs": len(doc.paragraphs),
        "characters": sum(len(p.text) for p in doc.paragraphs),
        "pageSizeCm": [round(section.page_width.cm, 1), round(section.page_height.cm, 1)],
        "fonts": sorted(fonts),
    }


def xlsx_facts(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(str(path))
    sheets = {}
    broken = 0
    for ws in wb.worksheets:
        formulas, numbers = 0, 0
        samples = []
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formulas += 1
                    broken += int("#REF!" in value)
                    if len(samples) < 3:
                        samples.append(f"{cell.coordinate}: {value[:80]}")
                elif isinstance(value, (int, float)):
                    numbers += 1
        sheets[ws.title] = {"rows": ws.max_row, "formulas": formulas, "numbers": numbers, "samples": samples}
    return {"sheets": sheets, "brokenRefs": broken}


def main() -> None:
    path = Path(sys.argv[1])
    kind = path.suffix.lower()
    facts = {".pptx": pptx_facts, ".docx": docx_facts, ".xlsx": xlsx_facts}.get(kind)
    if facts is None:
        print(json.dumps({"error": f"不认识的类型 {kind}"}, ensure_ascii=False))
        return
    try:
        print(json.dumps(facts(path), ensure_ascii=False))
    except Exception as error:  # noqa: BLE001 —— 打不开本身就是要记下的事实
        print(json.dumps({"error": f"打不开：{type(error).__name__}: {error}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
