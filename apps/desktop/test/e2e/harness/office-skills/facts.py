"""读出办公文件里机器能确定的事实，给 `office-skills.real.spec.mjs` 断言用。

    python facts.py <文件>

docx：按正文顺序的每一行（段落一行；表格每行一行，单元格用「 | 」连，合并单元格只留一份）
      加上页脚的文字；xlsx：每个非空单元格的值，公式原样（以 = 开头）、日期转成 ISO 字符串。
用办公扩展的解释器以 `-I` 跑：读的是被测 App 写出来的文件，不让工作区里的东西进 import 路径。
"""

import json
import re
import sys
import zipfile
from pathlib import Path


def docx_facts(path: Path) -> dict:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document._body).text.strip()
            if text:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, document._body).rows:
                cells: list[str] = []
                for cell in row.cells:
                    text = cell.text.strip()
                    if not cells or cells[-1] != text:
                        cells.append(text)
                lines.append(" | ".join(cells))
    with zipfile.ZipFile(path) as archive:
        footer_xml = " ".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer")
        )
    return {"lines": lines, "footer": re.sub(r"<[^>]+>", "", footer_xml)}


def xlsx_facts(path: Path) -> dict:
    import openpyxl

    book = openpyxl.load_workbook(str(path))
    cells = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                if hasattr(value, "isoformat"):
                    value = value.isoformat()
                elif not isinstance(value, (int, float, str)):
                    value = str(value)
                cells.append({"sheet": sheet.title, "ref": cell.coordinate, "value": value})
    return {"cells": cells}


def main() -> None:
    path = Path(sys.argv[1])
    try:
        facts = docx_facts(path) if path.suffix.lower() == ".docx" else xlsx_facts(path)
    except Exception as error:  # noqa: BLE001 —— 打不开本身就是要报告的事实
        facts = {"error": f"{type(error).__name__}: {error}"}
    print(json.dumps(facts, ensure_ascii=False))


if __name__ == "__main__":
    main()
