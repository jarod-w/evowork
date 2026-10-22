"""测试样本：用 openpyxl 直接造一份 xlsx，或把一份 xlsx 读成 JSON 给断言用。

造出来的是**用户手做的、或别的工具导出的那种文件**，不经过本技能的渲染器。
只拿渲染器自己的产物去测提取，测到的只是"渲染器的输出能读回来"——标题行、货币格式、
真日期、图表这些用户文件里最常见的东西一个都碰不到（2026-10-22 那几处就是这么漏的）。

用法：
  python3 workbook_fixture.py build <输出路径> <python 代码>
      代码里能直接用 wb（新工作簿）、ws（它的第一张表）、datetime、Image、png() 与
      openpyxl 的常用类；跑完自动保存。
  python3 workbook_fixture.py read <xlsx 路径>
      打印 {表名: {坐标: [值, data_type]}}；日期打印成 ISO 字符串。

两种都先 ensure_office_runtime：系统 python 没有 openpyxl（或插图要的 Pillow）时换到
办公扩展的解释器，与技能自己的跑法一致（见 test-helpers.ts 的 hasPythonModule）。
"""

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "_shared"))
from evowork_skill import ensure_office_runtime  # noqa: E402

ensure_office_runtime(("openpyxl", "PIL"))

import datetime  # noqa: E402
from io import BytesIO  # noqa: E402

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.drawing.image import Image  # noqa: E402
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402


def png() -> BytesIO:
    """一张 2×2 的图。openpyxl 插图要 Pillow —— 办公扩展里有（matplotlib 依赖它）。"""
    from PIL import Image as PILImage

    buffer = BytesIO()
    PILImage.new("RGB", (2, 2), "red").save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def build(out: str, code: str) -> None:
    wb = Workbook()
    namespace = {
        "wb": wb,
        "ws": wb.active,
        "datetime": datetime,
        "BarChart": BarChart,
        "Reference": Reference,
        "Comment": Comment,
        "Image": Image,
        "CellIsRule": CellIsRule,
        "ColorScaleRule": ColorScaleRule,
        "DataBarRule": DataBarRule,
        "DataValidation": DataValidation,
        "Font": Font,
        "PatternFill": PatternFill,
        "png": png,
    }
    exec(code, namespace)  # noqa: S102 —— 只跑测试文件里写的代码
    wb.save(out)


def read(path: str) -> None:
    book = load_workbook(path)
    found = {}
    for sheet in book.worksheets:
        cells = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                value = cell.value
                if isinstance(value, (datetime.date, datetime.time)):
                    value = value.isoformat()
                cells[cell.coordinate] = [value, cell.data_type]
        found[sheet.title] = cells
    print(json.dumps(found, ensure_ascii=False))


if __name__ == "__main__":
    if sys.argv[1] == "build":
        build(sys.argv[2], sys.argv[3])
    else:
        read(sys.argv[2])
