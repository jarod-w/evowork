#!/usr/bin/env python3
"""pptx / docx 的本机版式检查：不渲染，用字体度量与字形表查最常见的几类问题（2026-10-22）。

## 为什么要它

App 的版式预览要靠本机装了 LibreOffice（`services/ingest/src/office-preview.ts`），多数用户没有。
于是 agent 交付前的「视觉检查」只剩两条路：联网加载在线编辑器，或者现场用系统字体手写一个渲染器 ——
2026-10-12 的 PPT 任务两条都走了。它手写的那个查出来的问题恰好是能**不渲染就确定**的三类：
字体里没有的字显示成方框（那次是 U+2212 减号）、没指定字体的文字回落成别的字体、文字溢出文本框。
这里把它们做成确定的检查，离线、几百毫秒、不需要扩展之外的任何东西。

## 能说什么、不能说什么

- **溢出是估算**：按字体度量折行、行高按字号的 1.2 倍。差一两行的边缘情况会误判，所以只报
  超出可用高度半行以上的；它不替代在 Office 里看一眼。
- **缺字是确定的**：字体文件的字形表里没有这个码位，就一定显示不出来（或被系统替换成别的字体）。
- **本机没装的字体**（比如模板用的 Microsoft YaHei 在 Mac 上）只提示一次、按扩展自带的 Noto Sans SC
  估算；收件人的机器上多半有，不是要改的问题。

## 输出纪律

与四个技能一致（`evowork_skill.py` 头注释）：**不输出用户内容**。只报位置（第几页 / 哪个框 / 第几段）、
字号、码位与字体名。

用法：python3 layout_check.py <文件.pptx|文件.docx>   → 一行 JSON
"""

from __future__ import annotations

import json
import logging
import re
import sys
import unicodedata
from functools import lru_cache
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

EMU_PER_PT = 12700
LINE_SPACING = 1.2
DEFAULT_PT = 18.0
# python-pptx / PowerPoint 文本框的默认内边距：左右 0.1 英寸、上下 0.05 英寸
INSET_X_EMU = 91440
INSET_Y_EMU = 45720
OFF_SLIDE_TOLERANCE_EMU = 45720
# 量字用的放大倍数：PIL 的字号是整数像素，放大后再除回来，小字号也量得准
SCALE = 8
MAX_ISSUES = 30


def _issue(severity: str, kind: str, where: str, detail: str) -> dict[str, str]:
    return {"severity": severity, "kind": kind, "where": where, "detail": detail}


def _is_cjk(char: str) -> bool:
    return unicodedata.east_asian_width(char) in ("W", "F")


@lru_cache(maxsize=None)
def _fallback_font() -> Path | None:
    try:
        from evowork_skill import office_fonts_dir

        fonts = office_fonts_dir()
    except Exception:  # noqa: BLE001 —— 找不到扩展字体就退到 matplotlib 自带的
        fonts = None
    if fonts is not None:
        for candidate in sorted(fonts.glob("*.tt[fc]")) + sorted(fonts.glob("*.otf")):
            return candidate
    try:
        from matplotlib import font_manager

        return Path(font_manager.findfont(font_manager.FontProperties(family=["DejaVu Sans"])))
    except Exception:  # noqa: BLE001
        return None


@lru_cache(maxsize=None)
def _resolve_font(name: str) -> Path | None:
    """字体名 → 本机字体文件；没装就是 None（不回落到默认字体：那会把「没装」说成「装了」）。"""
    if not name:
        return None
    # findfont 找替代字重时会往 stderr 打一行告警；那不是这里要报的东西
    logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
    try:
        from matplotlib import font_manager

        return Path(
            font_manager.findfont(font_manager.FontProperties(family=[name]), fallback_to_default=False)
        )
    except Exception:  # noqa: BLE001 —— findfont 找不到时抛 ValueError
        return None


@lru_cache(maxsize=None)
def _face_index(path: Path, name: str) -> int:
    """.ttc 里与字体名对应的那一个字面；对不上就用第 0 个。"""
    if path.suffix.lower() != ".ttc":
        return 0
    try:
        from fontTools.ttLib import TTCollection

        for index, font in enumerate(TTCollection(str(path)).fonts):
            names = {record.toUnicode() for record in font["name"].names if record.nameID in (1, 4, 16)}
            if name in names:
                return index
    except Exception:  # noqa: BLE001
        pass
    return 0


@lru_cache(maxsize=None)
def _cmap(path: Path, name: str) -> frozenset[int]:
    try:
        from fontTools.ttLib import TTFont

        font = TTFont(str(path), fontNumber=_face_index(path, name), lazy=True)
        return frozenset(font.getBestCmap() or {})
    except Exception:  # noqa: BLE001
        return frozenset()


@lru_cache(maxsize=None)
def _pil_font(path: Path, name: str, size_pt: float):
    from PIL import ImageFont

    return ImageFont.truetype(str(path), max(1, round(size_pt * SCALE)), index=_face_index(path, name))


def _width_pt(text: str, font_path: Path, name: str, size_pt: float) -> float:
    return _pil_font(font_path, name, size_pt).getlength(text) / SCALE


TOKEN = re.compile(r"[A-Za-z0-9_.,:;%+\-/()'\"$€£¥#&@!?]+\s*|\s+|.", re.S)


def _lines_needed(text: str, width_pt: float, font_path: Path, name: str, size_pt: float) -> int:
    """贪心折行：西文按词、中文按字。空段也占一行。"""
    if width_pt <= 0:
        return 1
    lines, current = 1, 0.0
    for token in TOKEN.findall(text):
        token_width = _width_pt(token, font_path, name, size_pt)
        if current + token_width <= width_pt or current == 0:
            if token_width > width_pt and current == 0:
                # 一个词比整行还宽：按字强拆
                lines += int(token_width // width_pt)
                current = token_width % width_pt
            else:
                current += token_width
            continue
        lines += 1
        current = 0.0 if token.isspace() else token_width
    return lines


class _Fonts:
    """把一段文字用到的字体解析成文件，记下没装的与缺字的。"""

    def __init__(self) -> None:
        self.not_installed: set[str] = set()
        self.unspecified = 0

    def for_text(
        self, latin: str | None, east_asian: str | None, text: str, *, count: bool = True
    ) -> tuple[str, Path | None]:
        name = (east_asian if east_asian and any(_is_cjk(c) for c in text) else None) or latin or east_asian
        if not name:
            # 只按「文字片段」计一次；整段量高度时那次不算，否则同一处会被数两遍
            if count and text.strip():
                self.unspecified += 1
            return "", None
        path = _resolve_font(name)
        if path is None:
            self.not_installed.add(name)
        return name, path

    def measuring(self, name: str, path: Path | None) -> tuple[Path | None, str]:
        if path is not None:
            return path, name
        fallback = _fallback_font()
        return fallback, fallback.stem if fallback else ""


def _missing_glyphs(text: str, name: str, path: Path | None) -> list[str]:
    if path is None:
        return []
    cmap = _cmap(path, name)
    if not cmap:
        return []
    missing = {f"U+{ord(c):04X}" for c in text if not c.isspace() and unicodedata.category(c)[0] != "C" and ord(c) not in cmap}
    return sorted(missing)


# ───────────────────────────── pptx ─────────────────────────────


def _run_fonts(run, paragraph) -> tuple[str | None, str | None]:
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    latin = run.font.name
    ea_node = run._r.find(f".//{ns}ea")
    east_asian = ea_node.get("typeface") if ea_node is not None else None
    if not latin:
        latin = paragraph.font.name
    for typeface in (latin, east_asian):
        if typeface and typeface.startswith("+"):
            # 主题字体引用（+mn-lt / +mn-ea）：本机解析不了主题，按「没指定」处理
            return (None if latin and latin.startswith("+") else latin), None
    return latin, east_asian


def _size_pt(run, paragraph) -> float:
    size = run.font.size or paragraph.font.size
    return size.pt if size is not None else DEFAULT_PT


def _frame_issues(frame, box_width_emu: int, box_height_emu: int, where: str, fonts: _Fonts, issues: list) -> float:
    """检查一个文本框；返回估算的文字高度（pt），表格单元格要用它算行高。"""
    width_pt = max(0, box_width_emu - 2 * INSET_X_EMU) / EMU_PER_PT
    height_pt = max(0, box_height_emu - 2 * INSET_Y_EMU) / EMU_PER_PT
    needed_pt, line_pt = 0.0, DEFAULT_PT * LINE_SPACING
    for p_index, paragraph in enumerate(frame.paragraphs, start=1):
        runs = list(paragraph.runs) or [None]
        size = max((_size_pt(r, paragraph) for r in runs if r is not None), default=DEFAULT_PT)
        line_pt = size * LINE_SPACING
        text = "".join(r.text for r in paragraph.runs)
        latin, east_asian = _run_fonts(runs[0], paragraph) if runs[0] is not None else (paragraph.font.name, None)
        name, path = fonts.for_text(latin, east_asian, text, count=False)
        for run in paragraph.runs:
            r_latin, r_ea = _run_fonts(run, paragraph)
            r_name, r_path = fonts.for_text(r_latin, r_ea, run.text)
            missing = _missing_glyphs(run.text, r_name, r_path)
            if missing:
                issues.append(_issue("warn", "missing-glyph", f"{where}第 {p_index} 段", f"字体「{r_name}」里没有 {'、'.join(missing[:6])}，会显示成方框或被替换"))
        measure_path, measure_name = fonts.measuring(name, path)
        lines = _lines_needed(text, width_pt, measure_path, measure_name, size) if measure_path and text else 1
        needed_pt += lines * line_pt
    autofit = frame._txBody.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit") is not None
    if height_pt > 0 and not autofit and needed_pt > height_pt + line_pt / 2:
        issues.append(
            _issue("warn", "overflow", where, f"文字约需 {needed_pt:.0f}pt 高，框内只有 {height_pt:.0f}pt（估算，超出约 {(needed_pt - height_pt) / line_pt:.1f} 行）")
        )
    return needed_pt


def check_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation

    prs = Presentation(str(path))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    issues: list[dict[str, str]] = []
    fonts = _Fonts()
    for s_index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            label = f"第 {s_index} 页「{shape.name}」"
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            if None not in (left, top, width, height):
                if (
                    left < -OFF_SLIDE_TOLERANCE_EMU
                    or top < -OFF_SLIDE_TOLERANCE_EMU
                    or left + width > slide_w + OFF_SLIDE_TOLERANCE_EMU
                    or top + height > slide_h + OFF_SLIDE_TOLERANCE_EMU
                ):
                    issues.append(_issue("warn", "off-slide", label, "有一部分在页面之外"))
            if getattr(shape, "has_table", False) and shape.has_table and None not in (top, height):
                total = 0.0
                for r_index, row in enumerate(shape.table.rows):
                    tallest = 0.0
                    for c_index, cell in enumerate(row.cells):
                        col = shape.table.columns[c_index]
                        tallest = max(tallest, _frame_issues(cell.text_frame, col.width, 10**12, f"{label}表格第 {r_index + 1} 行第 {c_index + 1} 列", fonts, issues))
                    total += max(row.height / EMU_PER_PT if row.height else 0, tallest + 2 * INSET_Y_EMU / EMU_PER_PT)
                if top / EMU_PER_PT + total > slide_h / EMU_PER_PT + 4:
                    issues.append(_issue("warn", "table-off-slide", label, f"表格按内容撑开后约 {total:.0f}pt 高，会超出页面底部"))
            elif shape.has_text_frame and None not in (width, height):
                _frame_issues(shape.text_frame, width, height, label, fonts, issues)
    return _result(path, issues, fonts, slides=len(prs.slides))


# ───────────────────────────── docx ─────────────────────────────


def _heading_level(paragraph, w: str) -> int | None:
    """样式名（Heading N / 标题 N）或大纲级别 —— 随包 documents 技能用的是「正文 + 大纲级别」。"""
    style = paragraph.style.name if paragraph.style is not None else ""
    match = re.match(r"(?:Heading|标题)\s*(\d)", style)
    if match:
        return int(match.group(1))
    for element in (paragraph._p, paragraph.style.element if paragraph.style is not None else None):
        node = element.find(f".//{w}outlineLvl") if element is not None else None
        value = node.get(f"{w}val", "") if node is not None else ""
        if value.isdigit() and int(value) < 9:
            return int(value) + 1
    return None


def check_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    doc = Document(str(path))
    issues: list[dict[str, str]] = []
    fonts = _Fonts()
    section = doc.sections[0]
    page_w, page_h = section.page_width, section.page_height
    if page_w is not None and page_h is not None:
        if abs(page_w.cm - 21.59) < 0.2 and abs(page_h.cm - 27.94) < 0.2:
            issues.append(_issue("info", "paper", "第 1 节", "纸张是美式 Letter；国内办公一般用 A4"))
    text_width = None
    if page_w is not None and section.left_margin is not None and section.right_margin is not None:
        text_width = page_w - section.left_margin - section.right_margin
    last_level = 0
    for p_index, paragraph in enumerate(doc.paragraphs, start=1):
        current = _heading_level(paragraph, w)
        if current is not None:
            if last_level and current > last_level + 1:
                issues.append(_issue("warn", "heading-skip", f"第 {p_index} 段", f"标题从 {last_level} 级跳到 {current} 级"))
            last_level = current
        for run in paragraph.runs:
            rfonts = run._r.find(f".//{w}rFonts")
            east_asian = rfonts.get(f"{w}eastAsia") if rfonts is not None else None
            name, font_path = fonts.for_text(run.font.name, east_asian, run.text)
            missing = _missing_glyphs(run.text, name, font_path)
            if missing:
                issues.append(_issue("warn", "missing-glyph", f"第 {p_index} 段", f"字体「{name}」里没有 {'、'.join(missing[:6])}，会显示成方框或被替换"))
    if text_width is not None:
        for t_index, table in enumerate(doc.tables, start=1):
            widths = [column.width for column in table.columns if column.width is not None]
            if widths and sum(widths) > text_width * 1.02:
                issues.append(_issue("warn", "table-too-wide", f"第 {t_index} 张表", f"列宽合计 {sum(widths) / 360000:.1f}cm，超过版心 {text_width / 360000:.1f}cm"))
        for i_index, shape in enumerate(doc.inline_shapes, start=1):
            if shape.width is not None and shape.width > text_width * 1.02:
                issues.append(_issue("warn", "image-too-wide", f"第 {i_index} 张图", f"宽 {shape.width / 360000:.1f}cm，超过版心 {text_width / 360000:.1f}cm"))
    return _result(path, issues, fonts, tables=len(doc.tables))


def _result(path: Path, issues: list[dict[str, str]], fonts: _Fonts, **counts: int) -> dict[str, Any]:
    fallback = _fallback_font()
    notes = [
        _issue("info", "font-not-installed", "全文", f"本机没装「{name}」，按 {fallback.stem if fallback else '默认字体'} 估算；收件人的机器上多半有，不必改")
        for name in sorted(fonts.not_installed)
    ]
    if fonts.unspecified:
        notes.append(_issue("info", "font-unspecified", "全文", f"{fonts.unspecified} 处文字没指定字体，会跟随主题或系统默认（中文可能被替换）"))
    warnings = [i for i in issues if i["severity"] == "warn"]
    infos = [i for i in issues if i["severity"] == "info"] + notes
    return {
        "file": path.name,
        "warnings": warnings[:MAX_ISSUES],
        "notes": infos[:MAX_ISSUES],
        "truncated": max(0, len(warnings) - MAX_ISSUES),
        **counts,
    }


def check_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return check_pptx(path)
    if suffix == ".docx":
        return check_docx(path)
    return {"file": path.name, "warnings": [], "notes": [_issue("info", "unsupported", "全文", f"{suffix} 不做版式检查")], "truncated": 0}


def for_output(path: Path) -> dict[str, Any]:
    """渲染器成功输出里附带的检查结果。**检查本身出错不影响交付**：产物已经写好了，只记一条提示。"""
    try:
        result = check_file(path)
    except Exception as error:  # noqa: BLE001
        return {"layout_warnings": [], "layout_notes": [_issue("info", "check-failed", "全文", f"版式检查没跑成：{type(error).__name__}")]}
    return {"layout_warnings": result["warnings"], "layout_notes": result["notes"]}


def main() -> None:
    if len(sys.argv) != 2:
        print("用法：python3 layout_check.py <文件.pptx|文件.docx>", file=sys.stderr)
        sys.exit(2)
    try:
        from evowork_skill import ensure_office_runtime

        ensure_office_runtime(("pptx", "docx", "fontTools", "PIL", "matplotlib"))
    except ImportError:
        pass
    print(json.dumps(check_file(Path(sys.argv[1]).resolve()), ensure_ascii=False))


if __name__ == "__main__":
    main()
