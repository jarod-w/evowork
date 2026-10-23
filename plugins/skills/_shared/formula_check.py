"""生成工作簿的受限公式检查与缓存；不执行宏、外部链接或 Python 表达式。"""

from __future__ import annotations

import fnmatch
import io
import math
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from xml.etree import ElementTree as ET


class FormulaFault(Exception):
    pass


class UnverifiedFormula(Exception):
    pass


@dataclass(frozen=True)
class Area:
    sheet: str
    left: int
    top: int
    right: int
    bottom: int

    @property
    def size(self):
        return (self.right - self.left + 1) * (self.bottom - self.top + 1)


class FormulaCheck:
    def __init__(self, book):
        self.book = book
        self.sheets = {name.casefold(): name for name in book.sheetnames}
        self.memo = {}
        self.active = set()
        self.steps = 0
        self.depth = 0

    def step(self):
        self.steps += 1
        if self.steps > 2_000_000:
            raise UnverifiedFormula("calculation-limit")

    def area(self, token, here):
        from openpyxl.utils.cell import range_boundaries

        if "[" in token or "]" in token:
            raise UnverifiedFormula("external-reference")
        if "!" in token:
            name, reference = token.rsplit("!", 1)
            name = name[1:-1].replace("''", "'") if name.startswith("'") and name.endswith("'") else name
        else:
            name, reference = here, token
        sheet = self.sheets.get(name.casefold())
        if sheet is None:
            raise FormulaFault("#REF!")
        if not re.fullmatch(r"\$?[A-Za-z]{1,3}\$?[1-9]\d*(?::\$?[A-Za-z]{1,3}\$?[1-9]\d*)?", reference):
            raise UnverifiedFormula("reference-syntax")
        left, top, right, bottom = range_boundaries(reference)
        if left > right or top > bottom or right > 16384 or bottom > 1048576:
            raise FormulaFault("#REF!")
        area = Area(sheet, left, top, right, bottom)
        if area.size > 100_000:
            raise UnverifiedFormula("range-limit")
        return area

    def parse(self, formula, here):
        from openpyxl.formula import Tokenizer

        try:
            tokens = [token for token in Tokenizer(formula).items if token.type != "WHITE-SPACE"]
        except Exception:
            raise UnverifiedFormula("formula-syntax") from None
        position = 0
        precedence = {"=": 1, "<>": 1, "<": 1, ">": 1, "<=": 1, ">=": 1, "&": 2, "+": 3, "-": 3, "*": 4, "/": 4, "^": 5}

        def expression(minimum=0, depth=0):
            nonlocal position
            if depth > 100 or position >= len(tokens):
                raise UnverifiedFormula("formula-syntax")
            token = tokens[position]
            position += 1
            if token.type == "OPERAND":
                if token.subtype == "RANGE":
                    node = ("value", self.area(token.value, here))
                elif token.subtype == "NUMBER":
                    value = float(token.value)
                    if not math.isfinite(value):
                        raise FormulaFault("#NUM!")
                    node = ("value", value)
                elif token.subtype == "TEXT":
                    node = ("value", token.value[1:-1].replace('""', '"'))
                elif token.subtype == "LOGICAL":
                    node = ("value", token.value.upper() == "TRUE")
                elif token.subtype == "ERROR":
                    node = ("error", token.value)
                else:
                    raise UnverifiedFormula("operand-syntax")
            elif token.type == "OPERATOR-PREFIX" and token.value in ("+", "-"):
                node = ("unary", token.value, expression(6, depth + 1))
            elif token.type in ("FUNC", "PAREN") and token.subtype == "OPEN":
                function = token.type == "FUNC"
                args = []
                if position < len(tokens) and tokens[position].subtype != "CLOSE":
                    args.append(expression(0, depth + 1))
                    while function and position < len(tokens) and tokens[position].type == "SEP" and tokens[position].subtype == "ARG":
                        position += 1
                        args.append(expression(0, depth + 1))
                if position >= len(tokens) or tokens[position].type != token.type or tokens[position].subtype != "CLOSE":
                    raise UnverifiedFormula("formula-syntax")
                position += 1
                if function:
                    node = ("function", token.value[:-1].upper(), args)
                elif len(args) == 1:
                    node = args[0]
                else:
                    raise UnverifiedFormula("formula-syntax")
            else:
                raise UnverifiedFormula("formula-syntax")
            while position < len(tokens):
                following = tokens[position]
                if following.type == "OPERATOR-POSTFIX" and following.value == "%":
                    position += 1
                    node = ("unary", "%", node)
                    continue
                priority = precedence.get(following.value, -1) if following.type == "OPERATOR-INFIX" else -1
                if priority < minimum:
                    break
                position += 1
                node = ("binary", following.value, node, expression(priority + 1, depth + 1))
            return node

        tree = expression()
        if position != len(tokens):
            raise UnverifiedFormula("formula-syntax")
        return tree

    def cell(self, sheet, row, column):
        self.step()
        key = (sheet, row, column)
        if key in self.memo:
            return self.memo[key]
        if key in self.active:
            raise FormulaFault("#CIRC!")
        if self.depth >= 100:
            raise UnverifiedFormula("dependency-limit")
        worksheet = self.book[sheet]
        if row > worksheet.max_row or column > worksheet.max_column:
            return None
        cell = worksheet.cell(row, column)
        raw = cell.value
        if cell.data_type == "e":
            raise FormulaFault(str(raw))
        if cell.data_type == "f":
            self.active.add(key)
            self.depth += 1
            try:
                value = self.scalar(self.evaluate(self.parse(raw, sheet)))
                if value is not None and not isinstance(value, (str, bool, int, float)):
                    raise FormulaFault("#NUM!")
                if isinstance(value, (float, int)) and not math.isfinite(value):
                    raise FormulaFault("#NUM!")
            except ZeroDivisionError:
                raise FormulaFault("#DIV/0!") from None
            except (ValueError, ArithmeticError, TypeError):
                raise FormulaFault("#VALUE!") from None
            finally:
                self.depth -= 1
                self.active.discard(key)
        else:
            value = raw
            if isinstance(value, (date, datetime)):
                from openpyxl.utils.datetime import to_excel

                value = to_excel(value, self.book.epoch)
        self.memo[key] = value
        return value

    def values(self, area):
        return [self.cell(area.sheet, row, column) for row in range(area.top, area.bottom + 1) for column in range(area.left, area.right + 1)]

    def scalar(self, value):
        if isinstance(value, Area):
            if value.size != 1:
                raise UnverifiedFormula("array-expression")
            return self.cell(value.sheet, value.top, value.left)
        return value

    def number(self, value):
        value = self.scalar(value)
        if value is None or value == "":
            return 0.0
        try:
            return float(value)
        except (ValueError, TypeError):
            raise FormulaFault("#VALUE!") from None

    def compare(self, left, right, operator):
        left, right = self.scalar(left), self.scalar(right)
        if left is None:
            left = "" if isinstance(right, str) else 0
        if right is None:
            right = "" if isinstance(left, str) else 0
        if isinstance(left, bool) != isinstance(right, bool):
            raise UnverifiedFormula("mixed-comparison")
        if isinstance(left, str) and isinstance(right, str):
            left, right = left.casefold(), right.casefold()
        elif type(left) is not type(right) and not (isinstance(left, (int, float)) and isinstance(right, (int, float))):
            raise UnverifiedFormula("mixed-comparison")
        return {"=": lambda: left == right, "<>": lambda: left != right, ">": lambda: left > right, "<": lambda: left < right, ">=": lambda: left >= right, "<=": lambda: left <= right}[operator]()

    def truth(self, value):
        value = self.scalar(value)
        if isinstance(value, str):
            if value.upper() not in ("TRUE", "FALSE"):
                raise UnverifiedFormula("logical-coercion")
            return value.upper() == "TRUE"
        return bool(value)

    def evaluate(self, node):
        self.step()
        kind = node[0]
        if kind == "value":
            return node[1]
        if kind == "error":
            raise FormulaFault(node[1])
        if kind == "unary":
            value = self.number(self.evaluate(node[2]))
            return -value if node[1] == "-" else value / 100 if node[1] == "%" else value
        if kind == "binary":
            operator = node[1]
            left, right = self.evaluate(node[2]), self.evaluate(node[3])
            if operator in ("=", "<>", ">", "<", ">=", "<="):
                return self.compare(left, right, operator)
            if operator == "&":
                def text(value):
                    value = self.scalar(value)
                    if value is None:
                        return ""
                    if isinstance(value, bool):
                        return "TRUE" if value else "FALSE"
                    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
                return text(left) + text(right)
            left, right = self.number(left), self.number(right)
            if operator == "^" and (abs(right) > 1000 or abs(left) > 1e100):
                raise UnverifiedFormula("calculation-limit")
            return {"+": lambda: left + right, "-": lambda: left - right, "*": lambda: left * right, "/": lambda: left / right, "^": lambda: left ** right}[operator]()
        name, args = node[1:]
        if name == "IF" and 2 <= len(args) <= 3:
            return self.evaluate(args[1]) if self.truth(self.evaluate(args[0])) else self.evaluate(args[2]) if len(args) == 3 else False
        if name == "IFERROR" and len(args) == 2:
            try:
                return self.scalar(self.evaluate(args[0]))
            except FormulaFault as error:
                if str(error) == "#CIRC!":
                    raise
                return self.evaluate(args[1])
            except ZeroDivisionError:
                return self.evaluate(args[1])
        supported = {"SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA", "ROUND", "ABS", "LEN", "AND", "OR", "NOT", "SUMIF", "SUMIFS", "COUNTIF", "COUNTIFS", "AVERAGEIF", "AVERAGEIFS", "SUMPRODUCT"}
        if name not in supported:
            raise UnverifiedFormula("unsupported-function")
        values = [self.evaluate(arg) for arg in args]
        if name in {"SUMIF", "SUMIFS", "COUNTIF", "COUNTIFS", "AVERAGEIF", "AVERAGEIFS"}:
            return self.conditional(name, values)
        flattened, numeric = [], []
        for value in values:
            if isinstance(value, Area):
                items = self.values(value)
                flattened.extend(items)
                numeric.extend(item for item in items if isinstance(item, (int, float)) and not isinstance(item, bool))
            else:
                flattened.append(value)
                if value is not None and name in {"SUM", "AVERAGE", "MIN", "MAX", "COUNT"}:
                    if name == "COUNT":
                        try:
                            numeric.append(float(value))
                        except (ValueError, TypeError):
                            pass
                    else:
                        numeric.append(self.number(value))
        if name == "SUM":
            return sum(numeric)
        if name == "AVERAGE":
            return sum(numeric) / len(numeric)
        if name == "MIN":
            return min(numeric, default=0)
        if name == "MAX":
            return max(numeric, default=0)
        if name == "COUNT":
            return len(numeric)
        if name == "COUNTA":
            return sum(value is not None for value in flattened)
        if name == "ABS" and len(values) == 1:
            return abs(self.number(values[0]))
        if name == "ROUND" and len(values) == 2:
            digits = int(self.number(values[1]))
            if abs(digits) > 100:
                raise UnverifiedFormula("calculation-limit")
            return float(Decimal(str(self.number(values[0]))).quantize(Decimal(10) ** -digits, rounding=ROUND_HALF_UP))
        if name == "LEN" and len(values) == 1:
            value = self.scalar(values[0])
            if isinstance(value, bool):
                return len("TRUE" if value else "FALSE")
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            return len(str(value if value is not None else ""))
        if name == "AND":
            if any(not isinstance(value, (bool, int, float)) for value in flattened):
                raise UnverifiedFormula("logical-coercion")
            return all(flattened)
        if name == "OR":
            if any(not isinstance(value, (bool, int, float)) for value in flattened):
                raise UnverifiedFormula("logical-coercion")
            return any(flattened)
        if name == "NOT" and len(values) == 1:
            return not self.truth(values[0])
        if name == "SUMPRODUCT" and values:
            arrays = [self.values(value) if isinstance(value, Area) else [value] for value in values]
            shapes = {(value.bottom - value.top + 1, value.right - value.left + 1) if isinstance(value, Area) else (1, 1) for value in values}
            if len(shapes) != 1:
                raise FormulaFault("#VALUE!")
            return sum(math.prod(value if isinstance(value, (float, int)) and not isinstance(value, bool) else 0 for value in row) for row in zip(*arrays))
        raise UnverifiedFormula("function-arguments")

    def matches(self, value, criterion):
        criterion = self.scalar(criterion)
        if criterion is None:
            criterion = 0
        match = re.fullmatch(r"(<>|>=|<=|=|>|<)?(.*)", str(criterion), re.S)
        operator, operand = match.group(1) or "=", match.group(2)
        try:
            number = float(operand)
        except ValueError:
            number = None
        if number is not None:
            try:
                candidate = float(value) if value is not None else 0.0
            except (ValueError, TypeError):
                return operator == "<>"
            return self.compare(candidate, number, operator)
        if operator not in ("=", "<>"):
            return self.compare(str(value or ""), operand, operator)
        if "~" in operand:
            raise UnverifiedFormula("escaped-criterion")
        pattern = "".join("[[]" if char == "[" else "[]]" if char == "]" else char for char in operand.casefold())
        wildcard = "*" in operand or "?" in operand
        matched = (not wildcard or isinstance(value, str)) and fnmatch.fnmatchcase(str(value if value is not None else "").casefold(), pattern)
        return not matched if operator == "<>" else matched

    def conditional(self, name, values):
        plural = name.endswith("IFS")
        counting = name.startswith("COUNT")
        if plural:
            total, rest = (None, values) if counting else (values[0], values[1:]) if values else (None, [])
            if len(rest) < 2 or len(rest) % 2:
                raise UnverifiedFormula("function-arguments")
            pairs = list(zip(rest[::2], rest[1::2]))
        else:
            if len(values) not in ((2,) if counting else (2, 3)):
                raise UnverifiedFormula("function-arguments")
            pairs = [(values[0], values[1])]
            total = None if counting else values[2] if len(values) == 3 else values[0]
        ranges = [area for area, _ in pairs] + ([] if total is None else [total])
        if any(not isinstance(area, Area) for area in ranges):
            raise UnverifiedFormula("range-required")
        if len({(area.right - area.left, area.bottom - area.top) for area in ranges}) != 1:
            raise FormulaFault("#VALUE!")
        tests = [(self.values(area), criterion) for area, criterion in pairs]
        selected = [index for index in range(ranges[0].size) if all(self.matches(column[index], criterion) for column, criterion in tests)]
        if counting:
            return len(selected)
        width = total.right - total.left + 1
        amounts = [self.cell(total.sheet, total.top + index // width, total.left + index % width) for index in selected]
        numeric = [amount for amount in amounts if isinstance(amount, (int, float)) and not isinstance(amount, bool)]
        return sum(numeric) / len(numeric) if name.startswith("AVERAGE") else sum(numeric)


def verify_workbook(book):
    checker = FormulaCheck(book)
    issues, checked = [], {}
    formulas = [(sheet, cell) for sheet in book for row in sheet for cell in row if cell.data_type == "f"]
    for sheet, cell in formulas:
        try:
            checked[(sheet.title, cell.coordinate)] = checker.cell(sheet.title, cell.row, cell.column)
        except FormulaFault as error:
            issues.append({"kind": "formula-error", "sheet": book.sheetnames.index(sheet.title) + 1, "cell": cell.coordinate, "error": str(error)})
        except (UnverifiedFormula, RecursionError) as error:
            issues.append({"kind": "needs-review", "sheet": book.sheetnames.index(sheet.title) + 1, "cell": cell.coordinate, "reason": str(error) if isinstance(error, UnverifiedFormula) else "dependency-limit"})
        if len(issues) >= 30:
            break
    status = "failed" if any(issue["kind"] == "formula-error" for issue in issues) else "needs_review" if issues else "passed"
    return {"status": status, "checks": ["formula-references", "formula-values"], "formula_count": len(formulas), "checked_count": len(checked), "issues": issues}, checked


def save_with_cache(book, path, values):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not values:
        book.save(path)
        return
    original = io.BytesIO()
    book.save(original)
    result = io.BytesIO()
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    ET.register_namespace("", namespace[1:-1])
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
            if match:
                sheet = book.sheetnames[int(match.group(1)) - 1]
                root = ET.fromstring(payload)
                for cell in root.iter(namespace + "c"):
                    key = (sheet, cell.get("r"))
                    if cell.find(namespace + "f") is None or key not in values:
                        continue
                    value = values[key]
                    cached = cell.find(namespace + "v")
                    if cached is None:
                        cached = ET.SubElement(cell, namespace + "v")
                    if isinstance(value, str):
                        cell.set("t", "str")
                        cached.text = value
                    elif isinstance(value, bool):
                        cell.set("t", "b")
                        cached.text = "1" if value else "0"
                    else:
                        cell.set("t", "n")
                        cached.text = str(value if value is not None else 0)
                payload = ET.tostring(root, encoding="utf-8")
            target.writestr(info, payload)
    path.write_bytes(result.getvalue())
