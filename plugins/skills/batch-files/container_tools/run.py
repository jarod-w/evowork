#!/usr/bin/env python3
"""工作区内的批量文件操作：重命名、按扩展名归类、docx 转 PDF、合并 csv / xlsx。

四条保证，每条都有 `test/run.test.ts` 钉着：

1. **只在任务目录里动手。** `workspace` 是模型写进 JSON 的，可能是 `/` 或家目录。
   先取真实路径（realpath），它必须等于当前工作目录（内核跑命令的任务目录）的真实路径，
   或在它里面。系统自带的上层链接（macOS 的 /var → /private/var）因此不算数；
   工作区**里面**的符号链接一律拒绝，不跟随。
2. **不覆盖。** 同名按 NFC + casefold 判断 —— macOS 默认的文件系统不分大小写，
   `X.txt` 与 `x.txt`、`é` 的两种 Unicode 写法都是同一个文件，`rename` 会把前一个悄悄替换掉。
   真正落盘只用「建硬链接 + 删原名」：目标在检查之后被别的程序抢先占了，链接就失败，
   不会把它替换掉。临时文件用 O_EXCL 建，不跟随事先放好的符号链接。
3. **要么全做，要么不做。** 所有冲突（彼此撞名、撞已有文件、一个目标要当另一个目标的文件夹）
   在动手前查完。执行中途仍出错时，倒着撤回已经做的每一步。
4. **撤回也失败时如实报。** 退出码 6，消息逐条列出哪些改动已经生效、没撤回。
   任何情况下都不吐裸 traceback。

退出码与共用骨架一致（`plugins/skills/_shared/evowork_skill.py`），另加本技能独有的 6。
"""

from __future__ import annotations

import argparse
import csv
import errno
import io
import importlib.util
import os
import secrets
import stat
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO, Callable

SKILL_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SKILL_ROOT.parent / "_shared"))

from evowork_skill import (  # noqa: E402
    EXIT_INVALID_CONTENT,
    EXIT_PARTIAL,
    EXIT_RUNTIME_MISSING,
    ensure_office_runtime,
    fail,
    read_content,
    runtime_missing_message,
    succeeded,
    write_extracted,
)

MAX_FILES = 50
MAX_BYTES = 20 * 1024 * 1024

SYMLINK = "路径上有符号链接。本技能不跟随符号链接。"
SAME_NAME = "（名字只差大小写或 Unicode 写法也算同名）"


class Refused(Exception):
    """不能做：动手前查出来的，或执行中重新核对时发现的。"""

    def __init__(self, message: str, code: int = EXIT_INVALID_CONTENT) -> None:
        super().__init__(message)
        self.message = message
        self.code = code


def fold(name: str) -> str:
    """同名判断用的写法。macOS 默认不分大小写，也不分 Unicode 的组合写法（é 与 e + ́）。"""
    return unicodedata.normalize("NFC", unicodedata.normalize("NFC", name).casefold())


def fold_key(rel: Path) -> tuple:
    return tuple(fold(part) for part in rel.parts)


def show(rel: Path) -> str:
    return rel.as_posix()


def require_text(value: object, where: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise Refused(f"{where} 不能为空。")
    return value.strip()


def cap(count: int) -> None:
    if count > MAX_FILES:
        raise Refused(f"一次最多 {MAX_FILES} 个文件。拆开再跑，不要截断。")
    if count < 1:
        raise Refused("至少要有一个文件。")


def reject_unknown(raw: dict, allowed: set) -> None:
    if set(raw) - allowed:
        raise Refused("内容 JSON 有未声明的字段。")


# ── 路径：只认工作区里面、不经过符号链接的 ─────────────────────────────


@dataclass(frozen=True)
class Workspace:
    root: Path  # 真实路径
    given: Path  # JSON 里写的样子；绝对路径的条目可以用它当前缀


def open_workspace(raw: object) -> Workspace:
    text = require_text(raw, "workspace")
    given = Path(text)
    if not given.is_absolute():
        raise Refused("workspace 必须是绝对路径（不认 ~，也不认相对路径）。填当前任务目录。")
    if ".." in given.parts:
        raise Refused("workspace 不能含 ..。")
    root = Path(os.path.realpath(given))
    if not root.is_dir():
        raise Refused("workspace 必须是一个已经存在的文件夹。")
    task = Path(os.path.realpath(os.getcwd()))
    if root != task and task not in root.parents:
        raise Refused(
            "workspace 必须是当前任务目录，或它里面的文件夹。在任务目录里运行本脚本，workspace 填任务目录。"
        )
    return Workspace(root, given)


def lexical(ws: Workspace, raw: object, where: str) -> Path:
    """条目路径 → 相对工作区的路径。只做字面判断，不碰磁盘。"""
    text = require_text(raw, where)
    candidate = Path(text)
    if ".." in candidate.parts:
        raise Refused(f"{where}：路径不能跳出工作区（不能含 ..）。")
    rel = candidate
    if candidate.is_absolute():
        for base in (ws.given, ws.root):
            try:
                rel = candidate.relative_to(base)
                break
            except ValueError:
                continue
        else:
            raise Refused(f"{where}：路径超出工作区。")
    if not rel.parts:
        raise Refused(f"{where}：指向的是工作区本身，要给具体的文件。")
    return rel


def guard_chain(root: Path, rel: Path) -> None:
    """从工作区往下，rel 的每一段都不能是符号链接。不存在的部分到此为止。"""
    current = root
    for part in rel.parts:
        current = current / part
        try:
            info = os.lstat(current)
        except (FileNotFoundError, NotADirectoryError):
            return
        if stat.S_ISLNK(info.st_mode):
            raise Refused(SYMLINK)


def existing_file(ws: Workspace, rel: Path, where: str) -> os.stat_result:
    guard_chain(ws.root, rel)
    try:
        info = os.lstat(ws.root / rel)
    except (FileNotFoundError, NotADirectoryError):
        raise Refused(f"{where}：找不到要处理的文件。") from None
    if not stat.S_ISREG(info.st_mode):
        raise Refused(f"{where}：只处理普通文件，不处理文件夹。")
    if info.st_size > MAX_BYTES:
        raise Refused(f"文件超过 {MAX_BYTES} 字节。拆开再处理，不要截断。")
    return info


class Planner:
    """动手之前把所有来源和目标摆在一起查。

    查三件事：目标彼此不同名、不和磁盘上已有的东西同名、不拿另一个目标的位置当文件夹。
    同名一律按 `fold` 比较。已有的文件夹只差大小写时沿用它的写法，不在旁边再建一个。
    """

    def __init__(self, ws: Workspace) -> None:
        self.ws = ws
        self.sources: set = set()
        self.targets: set = set()
        self.new_dirs: dict = {}
        self._listings: dict = {}

    def add_source(self, rel: Path, duplicate_message: str) -> None:
        key = fold_key(rel)
        if key in self.sources:
            raise Refused(duplicate_message)
        self.sources.add(key)

    def _same_named(self, folder: Path, name: str) -> list:
        listing = self._listings.get(folder)
        if listing is None:
            listing = {}
            with os.scandir(self.ws.root / folder) as entries:
                for entry in entries:
                    listing.setdefault(fold(entry.name), []).append(entry.name)
            self._listings[folder] = listing
        return listing.get(fold(name), [])

    def resolve_target(
        self, rel: Path, *, source: os.stat_result | None, exists_message: str
    ) -> tuple:
        """返回（落盘用的相对路径, 要新建的文件夹, 目标是不是来源自己）。"""
        current = Path()
        dirs: list = []
        on_disk = True
        parts = rel.parts
        for index, part in enumerate(parts):
            key = fold_key(current / part)
            names = self._same_named(current, part) if on_disk else []
            if index == len(parts) - 1:
                if key in self.targets:
                    raise Refused(f"两个文件会落到同一个名字上{SAME_NAME}。不会覆盖。")
                if key in self.new_dirs:
                    raise Refused(f"一个目标和另一个目标要用的文件夹同名{SAME_NAME}。不会覆盖。")
                itself = False
                if names:
                    if source is None or len(names) != 1:
                        raise Refused(exists_message)
                    info = os.lstat(self.ws.root / current / names[0])
                    if (info.st_dev, info.st_ino) != (source.st_dev, source.st_ino):
                        raise Refused(exists_message)
                    itself = True
                self.targets.add(key)
                return current / part, dirs, itself
            if names:
                if len(names) > 1:
                    raise Refused("目标路径上有几个只差大小写的同名文件夹，分不清该用哪一个。")
                info = os.lstat(self.ws.root / current / names[0])
                if stat.S_ISLNK(info.st_mode):
                    raise Refused(SYMLINK)
                if not stat.S_ISDIR(info.st_mode):
                    raise Refused("目标路径上有一段已经是文件，不是文件夹。")
                current = current / names[0]
                continue
            on_disk = False
            if key in self.targets:
                raise Refused(f"一个目标要当文件夹用的路径，正是另一个目标文件{SAME_NAME}。")
            current = self.new_dirs.get(key, current / part)
            self.new_dirs[key] = current
            dirs.append(current)
        raise Refused("目标路径是空的。")


# ── 执行：每一步都记账，出错就倒着撤回 ─────────────────────────────


def link_no_replace(src: Path, dst: Path) -> None:
    """给 src 的内容建一个新名字 dst。dst 已存在（哪怕是悬空的符号链接）就抛 FileExistsError，绝不替换。"""
    if os.link in os.supports_follow_symlinks:
        os.link(src, dst, follow_symlinks=False)
    else:
        os.link(src, dst)


def file_mode() -> int:
    mask = os.umask(0)
    os.umask(mask)
    return 0o666 & ~mask


@dataclass
class Move:
    source: Path
    target: Path
    dirs: list
    case_only: bool = False  # 只改大小写 / Unicode 写法：不分大小写的盘上要经过一个中间名


@dataclass
class Journal:
    root: Path
    steps: list = field(default_factory=list)
    made: set = field(default_factory=set)

    def ensure_dirs(self, dirs: list) -> None:
        for rel in dirs:
            if rel in self.made:
                continue
            guard_chain(self.root, rel)
            os.mkdir(self.root / rel)
            self.made.add(rel)
            self.steps.append(("mkdir", rel))

    def move(self, src: Path, dst: Path) -> None:
        guard_chain(self.root, src)
        guard_chain(self.root, dst.parent)
        link_no_replace(self.root / src, self.root / dst)
        self.steps.append(("link", src, dst))
        os.unlink(self.root / src)
        self.steps[-1] = ("move", src, dst)

    def apply(self, op: Move) -> None:
        self.ensure_dirs(op.dirs)
        if op.source == op.target:
            return
        if op.case_only:
            middle = op.source.parent / f".evowork-rename-{secrets.token_hex(6)}"
            self.move(op.source, middle)
            self.move(middle, op.target)
        else:
            self.move(op.source, op.target)

    def create(self, dst: Path, write: Callable[[BinaryIO], object]) -> None:
        """写一个新文件：先写进同目录的临时文件（O_EXCL，不跟随符号链接），再不覆盖地挂到 dst 上。"""
        guard_chain(self.root, dst.parent)
        fd, temp = tempfile.mkstemp(prefix=".evowork-", suffix=".partial", dir=self.root / dst.parent)
        try:
            with os.fdopen(fd, "wb") as handle:
                if hasattr(os, "fchmod"):
                    os.fchmod(handle.fileno(), file_mode())
                write(handle)
            link_no_replace(Path(temp), self.root / dst)
            self.steps.append(("create", dst))
        finally:
            try:
                os.unlink(temp)
            except FileNotFoundError:
                pass

    def rollback(self) -> list:
        """倒着撤回。返回没能撤回的改动（按执行顺序）；空表 = 工作区已经恢复原样。"""
        left = []
        for step in reversed(self.steps):
            residue = self._undo(step)
            if residue:
                left.append(residue)
        left.reverse()
        return left

    def _undo(self, step: tuple) -> str | None:
        kind = step[0]
        try:
            if kind == "mkdir":
                os.rmdir(self.root / step[1])
                return None
            if kind == "create":
                os.unlink(self.root / step[1])
                return None
            if kind == "link":
                os.unlink(self.root / step[2])
                return None
            link_no_replace(self.root / step[2], self.root / step[1])
        except OSError as err:
            return f"{describe_step(step)}（没撤回：{reason(err)}）"
        try:
            os.unlink(self.root / step[2])
        except OSError as err:
            return (
                f"{show(step[1])} 已经恢复，但 {show(step[2])} 还是它的另一个名字（硬链接），没删掉：{reason(err)}"
            )
        return None


def describe_step(step: tuple) -> str:
    kind = step[0]
    if kind == "mkdir":
        return f"新建了文件夹 {show(step[1])}/"
    if kind == "create":
        return f"新建了文件 {show(step[1])}"
    if kind == "link":
        return f"{show(step[2])} 是 {show(step[1])} 的另一个名字（硬链接），{show(step[1])} 还在原处"
    return f"{show(step[1])} → {show(step[2])}"


def reason(err: BaseException) -> str:
    """出错原因的一句话。不带文件内容，只用错误的种类。"""
    if isinstance(err, Refused):
        return err.message
    if isinstance(err, FileExistsError):
        return "目标在执行时已经被别的程序占了，没有覆盖它"
    if isinstance(err, OSError):
        if err.errno == errno.EXDEV:
            return "来源和目标不在同一个磁盘卷上；本技能不做「复制再删除」式的移动"
        unsupported = {errno.EPERM, errno.EMLINK, getattr(errno, "ENOTSUP", -1), getattr(errno, "EOPNOTSUPP", -1)}
        if err.errno in unsupported:
            return "文件系统不允许这样移动（可能是不支持硬链接的磁盘，比如 exFAT / FAT 格式的 U 盘，或没有权限）"
        if err.errno == errno.EACCES:
            return "没有权限"
        return f"文件系统报错（{err.strerror or type(err).__name__}）"
    if isinstance(err, KeyboardInterrupt):
        return "被中断"
    return f"内部错误（{type(err).__name__}）"


def commit(ws: Workspace, out_path: Path, result: dict, apply: Callable[[Journal], None]) -> None:
    """唯一改动工作区的地方。连报告一起写完才算成功；任何一步出错都先撤回，再如实报。"""
    journal = Journal(ws.root)
    try:
        apply(journal)
        write_extracted(result, out_path)
    except BaseException as err:  # noqa: BLE001 —— 不论什么原因中断，都要先撤回
        cause = reason(err)
        left = journal.rollback()
        if not left:
            code = err.code if isinstance(err, Refused) else EXIT_INVALID_CONTENT
            fail(code, f"执行到一半出错：{cause}。已经做的改动都撤回了，工作区和执行前一样。")
        listing = "\n".join(f"- {line}" for line in left)
        fail(
            EXIT_PARTIAL,
            f"执行到一半出错：{cause}。撤回时又出错，下面这些改动已经生效、没有撤回"
            f"（工作区现在是半完成状态）。先把这些如实告诉用户，不要直接重跑：\n{listing}",
        )


# ── 四种操作 ─────────────────────────────


def moves(ops: list) -> Callable[[Journal], None]:
    def apply(journal: Journal) -> None:
        for op in ops:
            journal.apply(op)

    return apply


def rename_all(ws: Workspace, raw: dict, out_path: Path) -> None:
    items = raw.get("items")
    if not isinstance(items, list):
        raise Refused("items 必须是数组。")
    cap(len(items))
    planner = Planner(ws)
    parsed = []
    for index, item in enumerate(items):
        if not isinstance(item, dict) or set(item) - {"from", "to"}:
            raise Refused(f"items/{index} 只能有 from 和 to。")
        source = lexical(ws, item.get("from"), f"items/{index}/from")
        info = existing_file(ws, source, f"items/{index}/from")
        planner.add_source(source, f"重命名的来源重复{SAME_NAME}。")
        parsed.append((index, source, info, item.get("to")))
    ops = []
    for index, source, info, raw_to in parsed:
        wanted = lexical(ws, raw_to, f"items/{index}/to")
        target, dirs, itself = planner.resolve_target(
            wanted, source=info, exists_message=f"目标已经存在{SAME_NAME}。不会覆盖。"
        )
        if target == source:
            raise Refused("重命名的来源和目标是同一个文件。")
        ops.append(Move(source, target, dirs, case_only=itself))
    result = {
        "operation": "rename",
        "files": [{"from": show(op.source), "to": show(op.target)} for op in ops],
    }
    commit(ws, out_path, result, moves(ops))


def group_all(ws: Workspace, raw: dict, out_path: Path) -> None:
    files = raw.get("files")
    if not isinstance(files, list):
        raise Refused("files 必须是数组。")
    cap(len(files))
    planner = Planner(ws)
    parsed = []
    for index, item in enumerate(files):
        source = lexical(ws, item, f"files/{index}")
        info = existing_file(ws, source, f"files/{index}")
        planner.add_source(source, f"同一个文件列了两次{SAME_NAME}。")
        parsed.append((source, info))
    ops = []
    for source, info in parsed:
        folder_name = source.suffix.lower().lstrip(".") or "no-extension"
        if folder_name in {".", ".."} or "/" in folder_name or "\\" in folder_name:
            raise Refused("扩展名不能当作文件夹名。")
        target, dirs, itself = planner.resolve_target(
            Path(folder_name) / source.name,
            source=info,
            exists_message=f"归类后的位置上已经有同名文件{SAME_NAME}。不会覆盖。",
        )
        # 已经在该在的文件夹里：原地不动，也不顺手改它名字的大小写
        ops.append(Move(source, source if itself else target, dirs))
    result = {
        "operation": "group",
        "files": [{"from": show(op.source), "to": show(op.target)} for op in ops],
    }
    commit(ws, out_path, result, moves(ops))


def documents_pdf():
    path = SKILL_ROOT.parent / "documents" / "container_tools" / "render.py"
    spec = importlib.util.spec_from_file_location("evowork_documents_render", path)
    if spec is None or spec.loader is None:
        raise Refused("找不到文档技能的 PDF 转换。", EXIT_RUNTIME_MISSING)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def docx_to_pdf(ws: Workspace, raw: dict, out_path: Path) -> None:
    files = raw.get("files")
    if not isinstance(files, list):
        raise Refused("files 必须是数组。")
    cap(len(files))
    planner = Planner(ws)
    planned = []
    for index, item in enumerate(files):
        source = lexical(ws, item, f"files/{index}")
        if source.suffix.lower() != ".docx":
            raise Refused("docx-to-pdf 只接受 .docx。")
        existing_file(ws, source, f"files/{index}")
        planner.add_source(source, f"同一个 docx 列了两次{SAME_NAME}。")
        target, _dirs, _ = planner.resolve_target(
            source.with_suffix(".pdf"),
            source=None,
            exists_message=f"PDF 已经存在{SAME_NAME}。不会覆盖，docx 原件也还在。",
        )
        planned.append((source, target))
    module = documents_pdf()
    if module.office_converter() is None:
        raise Refused("需要本机 LibreOffice（soffice）才能把 docx 转成 PDF。安装后重试。", EXIT_RUNTIME_MISSING)
    with tempfile.TemporaryDirectory(prefix="evowork-batch-pdf-") as tmp:
        staged = []
        # 先全部转完、放在工作区外面。任何一份转不成（退出码 5）或这个平台不能隔离转换（3），
        # 转换器自己以那个退出码结束进程 —— 那时工作区里一个 PDF 都还没有。
        for source, target in planned:
            out = Path(tmp) / f"{len(staged)}.pdf"
            module.convert_docx_to_pdf(ws.root / source, out)
            staged.append((out, target))

        def apply(journal: Journal) -> None:
            for out, target in staged:
                journal.create(target, lambda handle, out=out: handle.write(out.read_bytes()))

        result = {"operation": "docx-to-pdf", "files": [{"pdf": show(target)} for _, target in staged]}
        commit(ws, out_path, result, apply)


def merged_csv(ws: Workspace, sources: list) -> bytes:
    rows: list = []
    header = None
    for source in sources:
        try:
            with (ws.root / source).open(newline="", encoding="utf-8") as handle:
                data = list(csv.reader(handle))
        except UnicodeDecodeError:
            raise Refused("有一份 csv 不是 UTF-8 编码。先另存为 UTF-8 再合并。") from None
        except csv.Error:
            raise Refused("有一份 csv 的格式读不了（引号或分隔符不成对）。") from None
        if not data or not any(any(cell.strip() for cell in row) for row in data):
            raise Refused("有一份 csv 是空的。")
        if header is None:
            header = data[0]
            rows.extend(data)
            continue
        if data[0] != header:
            raise Refused("csv 的表头不一致。不会把对不上的行并进去。")
        rows.extend(data[1:])
    buffer = io.StringIO(newline="")
    csv.writer(buffer).writerows(rows)
    return buffer.getvalue().encode("utf-8")


def merged_xlsx(ws: Workspace, sources: list):
    ensure_office_runtime(["openpyxl"])
    try:
        from openpyxl import Workbook, load_workbook
    except ModuleNotFoundError:
        raise Refused(runtime_missing_message("office", "合并 xlsx"), EXIT_RUNTIME_MISSING) from None
    book = Workbook()
    book.remove(book.active)
    used: set = set()
    for source in sources:
        try:
            loaded = load_workbook(ws.root / source)
        except Exception:  # noqa: BLE001 —— openpyxl 对坏文件抛的异常种类很多，原文不往外带
            raise Refused("有一份 xlsx 打不开（可能已损坏，或者其实不是 xlsx）。") from None
        for sheet in loaded.worksheets:
            title = sheet.title
            if title in used:
                title = f"{source.stem}-{sheet.title}"[:31]
            if title in used or any(ch in title for ch in ":\\/?*[]"):
                raise Refused("工作表重名或名字不能用。不会覆盖已有的表。")
            used.add(title)
            target = book.create_sheet(title)
            for row in sheet.iter_rows(values_only=True):
                target.append(list(row))
        loaded.close()
    if not book.worksheets:
        raise Refused("没有可合并的工作表。")
    return book


def merge_all(ws: Workspace, raw: dict, out_path: Path) -> None:
    sources_raw = raw.get("sources")
    if not isinstance(sources_raw, list):
        raise Refused("sources 必须是数组。")
    cap(len(sources_raw))
    if len(sources_raw) < 2:
        raise Refused("合并至少要两份文件。")
    sources = []
    for index, item in enumerate(sources_raw):
        source = lexical(ws, item, f"sources/{index}")
        existing_file(ws, source, f"sources/{index}")
        sources.append(source)
    suffixes = {path.suffix.lower() for path in sources}
    if suffixes == {".csv"}:
        kind = "csv"
    elif suffixes == {".xlsx"}:
        kind = "xlsx"
    else:
        raise Refused("一次只能合并同一种文件：全是 csv，或全是 xlsx。")
    wanted = lexical(ws, raw.get("out"), "out")
    if wanted.suffix.lower() != f".{kind}":
        raise Refused("合并结果的扩展名要和源文件一致。")
    target, dirs, _ = Planner(ws).resolve_target(
        wanted, source=None, exists_message=f"合并结果不能覆盖已有文件，也不能写回源文件{SAME_NAME}。"
    )
    # 先在内存里合好：表头不一致、文件打不开这类问题都在动手之前暴露
    if kind == "csv":
        data = merged_csv(ws, sources)

        def write(handle: BinaryIO) -> None:
            handle.write(data)

    else:
        book = merged_xlsx(ws, sources)

        def write(handle: BinaryIO) -> None:
            book.save(handle)

    def apply(journal: Journal) -> None:
        journal.ensure_dirs(dirs)
        journal.create(target, write)

    commit(ws, out_path, {"operation": "merge", "out": show(target), "sources": len(sources)}, apply)


OPERATIONS = {
    "rename": ({"operation", "workspace", "items"}, rename_all),
    "group": ({"operation", "workspace", "files"}, group_all),
    "docx-to-pdf": ({"operation", "workspace", "files"}, docx_to_pdf),
    "merge": ({"operation", "workspace", "sources", "out"}, merge_all),
}


def main() -> None:
    parser = argparse.ArgumentParser(description="工作区内批量重命名、归类、转 PDF 或合并表格")
    parser.add_argument("--content", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    out_path = Path(args.out).resolve()
    if out_path.suffix.lower() != ".json":
        fail(EXIT_INVALID_CONTENT, "--out 必须是 .json。")
    raw = read_content(Path(args.content).resolve())
    operation = raw.get("operation")
    try:
        ws = open_workspace(raw.get("workspace"))
        entry = OPERATIONS.get(operation) if isinstance(operation, str) else None
        if entry is None:
            raise Refused("operation 只能是 rename、group、docx-to-pdf 或 merge。")
        allowed, run = entry
        reject_unknown(raw, allowed)
        run(ws, raw, out_path)
    except Refused as err:
        fail(err.code, err.message)
    except SystemExit:
        raise
    except BaseException as err:  # noqa: BLE001 —— 不吐裸 traceback。改动只发生在 commit 里，它自己撤回并报错
        fail(EXIT_INVALID_CONTENT, f"出错了，还没有改动任何文件：{reason(err)}。")
    succeeded(out_path, operation=operation)


if __name__ == "__main__":
    main()
